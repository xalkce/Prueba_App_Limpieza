import datetime
import io
import traceback
import urllib.parse
import openpyxl
import pandas as pd
import streamlit as st

# =============================================================
# PROCESADOR DE FLUJOS DE CAJA - KCE
# Copyright (c) 2026 Kaizaharra Corporación Empresarial (KCE).
# Todos los derechos reservados.
#
# AVISO DE CONFIDENCIALIDAD Y PROPIEDAD INTELECTUAL:
# Este software y su código fuente son propiedad exclusiva y confidencial de KCE. 
# Queda terminantemente prohibida su copia, reproducción, cesión, plagio, distribución 
# total o parcial a terceros ajenos a la organización sin autorización expresa.
# =============================================================

# Configuracion Basica de la Página
st.set_page_config(
    page_title="Procesador de Flujos de Caja - KCE",
    page_icon="📊",
    layout="wide",
)

# Boton de Soporte Arriba a la Derecha
col_titulo, col_ayuda = st.columns([4, 1.2])

with col_titulo:
    st.title("📊 Procesador y Limpieza de Cuentas")

# Pop Up de Asistencia General / Feedback
@st.dialog("💬 Asistencia y Soporte")
def popup_soporte():
    st.markdown("Describe qué no funciona, qué dato no cuadra o qué sugerencia de diseño tienes:")
    comentario = st.text_area("Detalle de la incidencia o sugerencia:", placeholder="Ej: No me convence el formato de las fechas o el proyecto X no aparece...", height=130,)

    email_soporte = "xalmodovar@kce.es"
    asunto_soporte = urllib.parse.quote("Consulta / Incidencia en Procesador de Flujos KCE")
    cuerpo_soporte = urllib.parse.quote( f"Hola Xabier,\n\nTe contacto por la siguiente incidencia o sugerencia en la aplicación:\n\n{comentario}\n")
    url_soporte = (f"mailto:{email_soporte}?subject={asunto_soporte}&body={cuerpo_soporte}")

    st.link_button(label=f"📧 Enviar correo a {email_soporte}", url=url_soporte, type="primary", use_container_width=True,)


with col_ayuda:
    st.write("")  # Espaciador vertical
    if st.button( "¿Algún problema / sugerencia?", use_container_width=True, type="secondary"):
        popup_soporte()

st.markdown(
    """
Sube el archivo Excel de valoraciones (`.xlsx`). La aplicación extraerá las transacciones, 
limpiará los metadatos y te permitirá consultar métricas de cartera y descargar el **CSV**.
"""
)


# Pop Up por Error
@st.dialog("⚠️ Error al Procesar el Archivo")
def popup_error(error_msg: str, detalle_tb: str):
    st.error("Se ha producido un error durante la lectura o transformación del Excel.")

    with st.expander("Ver detalle técnico del error"):
        st.code(detalle_tb if detalle_tb else error_msg, language="python")

    st.markdown("Haz clic en el siguiente botón para generar un correo de aviso automático a soporte técnico:")

    # Mail To
    email_soporte = "xalmodovar@kce.es"
    asunto = urllib.parse.quote("Aviso de Incidencia: Error en Procesador KCE")
    cuerpo = urllib.parse.quote(
        f"Hola Xabier,\n\n"
        f"Se ha producido un error al procesar el archivo Excel en la aplicación:\n\n"
        f"Error: {error_msg}\n\n"
        f"Detalle técnico:\n{detalle_tb[:600]}...\n\n"
        f"Por favor, revisa el archivo o el código de la app."
    )
    url_mailto = f"mailto:{email_soporte}?subject={asunto}&body={cuerpo}"

    st.link_button(
        label=f"📧 Enviar aviso por correo a {email_soporte}",
        url=url_mailto,
        type="primary",
        use_container_width=True,
    )


# Extraccion y Limpieza


def procesar_excel(archivo_buffer):
    wb = openpyxl.load_workbook(archivo_buffer, data_only=True)
    transacciones_consolidadas = []

    for sheet in wb.sheetnames:
        # Sep. Vacios
        if "-->" in sheet:
            continue

        # 44 - duplicado
        if sheet.strip().lower() == "proyecto 44-préstamo":
            continue

        ws = wb[sheet]

        # Info Basica
        meta = {"sheet_name": sheet}
        for r in range(4, 12):
            clave = ws.cell(r, 2).value
            valor = ws.cell(r, 3).value
            if clave:
                meta[str(clave).strip()] = (
                    str(valor).strip() if valor is not None else ""
                )

        # Nombre en el 27
        if sheet == "Proyecto 27":
            meta["Proyecto"] = "Proyecto 27"

        # Fecha
        header_r = None
        for r in range(12, 18):
            if ws.cell(r, 2).value == "Fecha":
                header_r = r
                break

        if not header_r:
            continue

        # Valoracion Neta en Alquileres
        es_desglose_neto = ws.cell(header_r, 4).value == "Valoración neta"

        # Transacciones
        for r in range(header_r + 1, ws.max_row + 1):
            fecha_val = ws.cell(r, 2).value
            col3_val = ws.cell(r, 3).value
            col4_val = ws.cell(r, 4).value

            if es_desglose_neto:
                importe_val = col4_val
                grupo_val = ws.cell(r, 5).value
                concepto_val = ws.cell(r, 6).value
            else:
                importe_val = col3_val
                grupo_val = col4_val
                concepto_val = ws.cell(r, 5).value

            # Celdas Vacias
            if fecha_val is None or importe_val is None:
                continue

            # Normalizar fechas
            if isinstance(fecha_val, (datetime.datetime, datetime.date)):
                fecha = pd.to_datetime(fecha_val)
            else:
                fecha = pd.to_datetime(
                    str(fecha_val).strip(), dayfirst=True, errors="coerce"
                )

            if pd.isna(fecha):
                continue

            # Float con dos decimales
            try:
                importe_final = round(float(importe_val), 2)
            except (ValueError, TypeError):
                importe_final = 0.0

            # Grupo y Concepto
            grupo = str(grupo_val).strip() if grupo_val is not None else ""
            concepto = (
                str(concepto_val).strip() if concepto_val is not None else ""
            )

            if grupo.lower() in ["dividendo", "dividendos"]:
                grupo = "Dividendos"
            elif grupo.lower() in ["inversión", "inversion"]:
                grupo = "Inversión"
            elif grupo.lower() in ["préstamo", "prestamo"]:
                grupo = "Préstamo"

            transacciones_consolidadas.append(
                {
                    "Proyecto": meta.get("Proyecto", sheet),
                    "Categoria": meta.get("Categoría", ""),
                    "Subcategoria": meta.get("Subcategoría", ""),
                    "Tipo": meta.get("Tipo", ""),
                    "Partner": meta.get("Partner", "Sin"),
                    "Status": meta.get("Status", "Abierto"),
                    "Fecha": fecha.strftime("%Y-%m-%d"),
                    "Importe (€)": importe_final,
                    "Grupo": grupo,
                    "Concepto": concepto,
                }
            )

    columnas_ordenadas = [
        "Proyecto",
        "Categoria",
        "Subcategoria",
        "Fecha",
        "Importe (€)",
        "Tipo",
        "Concepto",
        "Status",
        "Partner",
    ]

    if not transacciones_consolidadas:
        return pd.DataFrame(columns=columnas_ordenadas)

    df = pd.DataFrame(transacciones_consolidadas)
    return df[columnas_ordenadas]


# Flujo Principal
archivo_subido = st.file_uploader(
    "Selecciona o arrastra el archivo Excel (.xlsx)",
    type=["xlsx"],
    help="Debe ser el archivo original de valoraciones.",
)

if archivo_subido is not None:
    with st.spinner("Procesando pestañas y calculando métricas..."):
        try:
            df_resultado = procesar_excel(archivo_subido)

            if df_resultado.empty:
                st.warning(
                    "No se encontraron transacciones válidas en el archivo subido."
                )
            else:
                st.success(
                    f"✅ Procesamiento completado: **{len(df_resultado):,}** transacciones registradas."
                )

                # Metricas
                df_proyectos_unicos = df_resultado.drop_duplicates(
                    subset=["Proyecto"]
                )
                total_proyectos = len(df_proyectos_unicos)

                proyectos_abiertos = (
                    df_proyectos_unicos["Status"]
                    .str.strip()
                    .str.lower()
                    .eq("abierto")
                    .sum()
                )
                proyectos_cerrados = (
                    df_proyectos_unicos["Status"]
                    .str.strip()
                    .str.lower()
                    .eq("cerrado")
                    .sum()
                )

                st.subheader("📌 Resumen de Cartera")
                m1, m2, m3, m4 = st.columns(4)

                # Total de Movimientos
                m1.metric("Total Movimientos", f"{len(df_resultado):,}")

                # proyectos abiertos
                m2.metric(
                    "Proyectos Abiertos",
                    f"{proyectos_abiertos}/{total_proyectos}",
                    help=f"Hay {proyectos_abiertos} proyectos abiertos de un total de {total_proyectos}",
                )

                # proyectos cerrados
                m3.metric(
                    "Proyectos Cerrados",
                    f"{proyectos_cerrados}/{total_proyectos}",
                )

                # flujos totales
                m4.metric(
                    "Total Flujos (€)",
                    f"{df_resultado['Importe (€)'].sum():,.2f} €",
                )

                st.divider()

                st.subheader("💰 Flujos de Caja Acumulados")

                tab_empresa, tab_partner = st.tabs(
                    ["🏢 Por Empresa / Proyecto", "🤝 Por Partner"]
                )

                # flujos por empresa
                with tab_empresa:
                    df_por_empresa = (
                        df_resultado.groupby("Proyecto", as_index=False)[
                            "Importe (€)"
                        ]
                        .sum()
                        .sort_values(by="Importe (€)", ascending=False)
                    )

                    c_left, c_right = st.columns([1.2, 1])
                    with c_left:
                        st.dataframe(
                            df_por_empresa,
                            column_config={
                                "Importe (€)": st.column_config.NumberColumn(
                                    "Suma Flujos (€)",
                                    format="%.2f €",
                                )
                            },
                            use_container_width=True,
                            height=350,
                        )
                    with c_right:
                        st.bar_chart(
                            df_por_empresa.set_index("Proyecto")[
                                "Importe (€)"
                            ],
                            height=350,
                        )
                # por partner
                with tab_partner:
                    df_por_partner = (
                        df_resultado.groupby("Partner", as_index=False)[
                            "Importe (€)"
                        ]
                        .sum()
                        .sort_values(by="Importe (€)", ascending=False)
                    )

                    c_left_p, c_right_p = st.columns([1.2, 1])
                    with c_left_p:
                        st.dataframe(
                            df_por_partner,
                            column_config={
                                "Importe (€)": st.column_config.NumberColumn(
                                    "Suma Flujos (€)",
                                    format="%.2f €",
                                )
                            },
                            use_container_width=True,
                            height=350,
                        )
                    with c_right_p:
                        st.bar_chart(
                            df_por_partner.set_index("Partner")[
                                "Importe (€)"
                            ],
                            height=350,
                        )

                st.divider()

                # Descargar el CSV
                st.subheader("📄 Transacciones Consolidadas")
                st.dataframe(df_resultado, use_container_width=True, height=350)

                csv_bytes = df_resultado.to_csv(
                    index=False, encoding="utf-8-sig"
                ).encode("utf-8-sig")

                st.download_button(
                    label="📥 Descargar Transacciones_Totales.csv",
                    data=csv_bytes,
                    file_name="Transacciones_Totales.csv",
                    mime="text/csv",
                    type="primary",
                )

        # para mostrar el Pop-Up en caso de error
        except Exception as e:
            tb_str = traceback.format_exc()
            popup_error(str(e), tb_str)
            st.error(f"❌ Error al procesar el archivo: {e}")

# Aviso Legal de Propiedad
st.markdown("<br><br>", unsafe_allow_html=True)
st.divider()

st.caption(
    """
    **© 2026 Kaizaharra Corporación Empresarial (KCE). Todos los derechos reservados.**  
    *Herramienta interna y confidencial. Queda prohibida la reproducción, copia, distribución, 
    modificación o ingeniería inversa de este software y sus algoritmos de transformación de datos.*
    """
)
